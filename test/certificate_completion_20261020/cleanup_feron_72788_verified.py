from __future__ import annotations
import datetime, hashlib, json, sys
from pathlib import Path
from openpyxl import load_workbook
sys.path[:0]=['/opt/homebrew/lib/python3.14/site-packages','test/certificate_api_pilot_20261012']
from api_client import Client
from reconcile_approved_ozon_docs_to_wb import doc_key, wb_datetime
from run_verified_feron_group_20261020 import wb_status
ROOT=Path('/Users/vladimirgrebennikov/Downloads/Новая папка/Документы_качества_Feron')
WORKBOOK=ROOT/'Сопоставление_артикулов_и_файлов_документов.xlsx'
ORIGINAL=ROOT/'По_артикулам/41718-1/02_Декларация.jpg'
STAGED=Path('/Users/vladimirgrebennikov/Code/Checksheets_Project/Checksheets/test/certificate_completion_20261020/source_docs/feron_72788_23_upload.pdf')
RECEIPT=Path('/Users/vladimirgrebennikov/Code/Checksheets_Project/Checksheets/test/certificate_completion_20261020/feron_72788_local_cleanup_receipt_20261020.json')
NUMBER='ЕАЭС N RU Д-CN.РА04.В.72788/23'; CID=5910474; OFFER='41718-1'; PID='674855094'; SKU='1242302791'; ISSUE='2023-06-20'; EXPIRY='2028-06-19'
assert ORIGINAL.is_file() and STAGED.is_file()
orig_hash=hashlib.sha256(ORIGINAL.read_bytes()).hexdigest(); staged_hash=hashlib.sha256(STAGED.read_bytes()).hexdigest()
client=Client()
meta=(client.call('ozon','/v1/product/certificate/info',{'certificate_number':NUMBER}).get('result') or {})
assert str(meta.get('certificate_id'))==str(CID) and meta.get('status_code')=='approved',meta
bindings=client.ozon_bindings(CID)
exact=[x for x in bindings if str(x.get('product_id'))==PID and str(x.get('sku'))==SKU]
assert len(exact)==1 and exact[0].get('product_status_code')=='approved',exact
cards=client.wb_cards(OFFER)
assert len(cards)==1 and cards[0].get('vendorCode')==OFFER and str(cards[0].get('brand','')).casefold()=='feron'
wbs=wb_status(cards[0],NUMBER,2,ISSUE,EXPIRY)
assert wbs['status']=='valid',wbs
wb=load_workbook(WORKBOOK,read_only=True,data_only=True); sh=wb['Артикул-файл']; headers=[c.value for c in sh[5]]; ix={x:i for i,x in enumerate(headers)}
refs=[]
for r in sh.iter_rows(min_row=6,values_only=True):
 if r and r[ix['Файл в папке артикула']]=='По_артикулам/41718-1/02_Декларация.jpg': refs.append({'offer_id':str(r[ix['Артикул продавца']]),'number':str(r[ix['Номер документа / реквизиты']])})
assert refs==[{'offer_id':OFFER,'number':'72788/23'}],refs
result={'checked_at':datetime.datetime.now(datetime.timezone.utc).isoformat(),'certificate_number':NUMBER,'ozon_certificate_id':CID,'ozon_certificate_status':meta['status_code'],'ozon_binding':{'offer_id':OFFER,'product_id':PID,'sku':SKU,'status':exact[0]['product_status_code']},'wb_status':wbs['status'],'wb_document_visible':wbs['document_visible'],'wb_dates_match':wbs['dates_match'],'source_workbook_references':refs,'deleted_files':[]}
for p,h in ((ORIGINAL,orig_hash),(STAGED,staged_hash)):
 assert hashlib.sha256(p.read_bytes()).hexdigest()==h
 size=p.stat().st_size
 p.unlink()
 result['deleted_files'].append({'path':str(p),'sha256':h,'bytes':size})
for x in result['deleted_files']:
 x['deleted']=not Path(x['path']).exists()
RECEIPT.write_text(json.dumps(result,ensure_ascii=False,indent=2))
print(json.dumps(result,ensure_ascii=False,indent=2))
