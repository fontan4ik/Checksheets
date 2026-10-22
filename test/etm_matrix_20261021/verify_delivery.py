import json,shutil,hashlib
from pathlib import Path
from collections import Counter
import openpyxl
here=Path(__file__).parent
d=json.loads((here/'matched.json').read_text())
source=here.parents[1]/'outputs/etm_matrix_20261021/Сопоставление_Ozon_ЭТМ_21.10.2026.xlsx'
w=openpyxl.load_workbook(source,read_only=True,data_only=True)
s=w['Матрица']; headers=next(s.iter_rows(min_row=8,max_row=8,values_only=True));assert list(headers)==d['headers']
rows=list(s.iter_rows(min_row=9,values_only=True));assert len(rows)==d['output_rows']
expected={(r[13],r[2]):r for r in d['rows']}
assert len({r[13] for r in rows})==d['total']
for r in rows:
    key=(r[13],r[2] or '')
    assert list(r)==[None if v=='' else v for v in expected[key]],key
    assert type(r[5]) is int and r[5] in (0,1)
assert len({r[13] for r in rows if r[1] is None})==63
assert Counter(r[5] for r in rows)==Counter(r[5] for r in d['rows'])
brandrows=list(w['По брендам'].iter_rows(min_row=9,max_row=8+len(d['brand_rows']),max_col=6,values_only=True))
assert [list(r) for r in brandrows]==[[None if v=='' else v for v in r] for r in d['brand_rows']]
w.close()
dest=Path('/Users/vladimirgrebennikov/Downloads')/source.name
if dest.exists():shutil.copy2(dest,here/'delivered_before_revision.xlsx')
shutil.copy2(source,dest)
assert hashlib.sha256(source.read_bytes()).digest()==hashlib.sha256(dest.read_bytes()).digest()
print(json.dumps(dict(path=str(dest),unique_products=d['total'],rows=len(rows),bound_products=d['bound_products'],problems=d['problem_products']),ensure_ascii=False))
