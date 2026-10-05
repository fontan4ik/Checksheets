import json
import re
import threading
import time
from concurrent.futures import ThreadPoolExecutor, as_completed
from decimal import Decimal, InvalidOperation
from pathlib import Path
from urllib.parse import quote

from openpyxl import load_workbook

import config
from etm_sync_multi_store import create_etm_session, login_etm
from rs_sync_local import create_rs_session, get_rs_headers


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "outputs" / "01a0f179-fa02-7ac1-8c59-51bb8848e624"
SOURCE_INSPECT = OUT_DIR / "DKC_остатки_по_складам.xlsx.inspect.ndjson"
RS_FILES = [OUT_DIR / "Русский_свет_Самара.xlsx", OUT_DIR / "Русский_свет_Москва.xlsx"]
CHECKPOINT = Path("/tmp/dkc_package_fetch_progress.json")
RESULT = Path("/tmp/dkc_package_fetch_result.json")
MIN_START_GAP = 1.0
PRINT_EVERY = 60


def norm(value):
    return re.sub(r"[^A-Z0-9]", "", str(value or "").upper())


def load_rows():
    table = None
    with SOURCE_INSPECT.open(encoding="utf-8") as source:
        for line in source:
            item = json.loads(line)
            if item.get("kind") == "table" and item.get("sheet") == "DKC остатки":
                table = item["values"]
                break
    if table is None:
        raise RuntimeError("Не найден лист с остатками DKC")
    header_index = next(i for i, row in enumerate(table) if row and row[0] == "Артикул DKC")
    headers = table[header_index]
    rows = [dict(zip(headers, row)) for row in table[header_index + 1 :] if row and row[0]]
    return rows


def load_rs_codes(articles):
    codes_by_article = {}
    wanted = set(articles)
    for path in RS_FILES:
        print(f"Читаю каталог Русского Света: {path.name}", flush=True)
        book = load_workbook(path, read_only=True, data_only=True)
        sheet = book.active
        for row in sheet.iter_rows(min_row=7, values_only=True):
            if len(row) < 5 or str(row[4] or "").strip().casefold() != "dkc":
                continue
            article_key = norm(row[0])
            rs_code = str(row[2] or "").strip()
            if article_key in wanted and rs_code:
                codes_by_article.setdefault(article_key, set()).add(rs_code)
        book.close()
    return codes_by_article


def num(value):
    if value in (None, ""):
        return None
    try:
        parsed = Decimal(str(value).replace(",", ".").strip())
    except (InvalidOperation, ValueError):
        return None
    if not parsed.is_finite() or parsed <= 0:
        return None
    return parsed


def show_num(value):
    if value is None:
        return ""
    if value == value.to_integral_value():
        return str(int(value))
    return format(value.normalize(), "f").replace(".", ",")


def format_package(value, unit):
    n = num(value)
    if n is None:
        return ""
    unit = str(unit or "").strip().replace(".", "")
    unit = re.sub(r"\s+", " ", unit)
    return f"{show_num(n)} {unit}".strip()


def choose_etm_package(data):
    unit = data.get("gdsUnitName") or ""
    packs = data.get("gdsPacks") or []
    if isinstance(packs, list):
        by_code = {str(p.get("gdsPackCode")): p for p in packs if isinstance(p, dict)}
        # Pack level 3 is the manufacturer package shown in the ETM catalog.
        for code in ("3", "2"):
            pack = by_code.get(code)
            value = num(pack.get("gdsPackVal")) if pack else None
            if value is not None and value > 1:
                return format_package(value, unit), f"ETM gdsPacks[{code}]"
    min_pack = num(data.get("minPack"))
    if min_pack is not None:
        return format_package(min_pack, unit), "ETM minPack"
    if isinstance(packs, list):
        unit_pack = next(
            (p for p in packs if isinstance(p, dict) and str(p.get("gdsPackCode")) == "1"),
            None,
        )
        if unit_pack:
            return format_package(unit_pack.get("gdsPackVal"), unit_pack.get("gdsPackName") or unit), "ETM gdsPacks[1]"
    return "", ""


def choose_rs_package(info):
    if not isinstance(info, dict):
        return "", ""
    primary_uom = str(info.get("PRIMARY_UOM") or "").strip()
    # RS can encode factory package size directly as "УП.100 м".
    match = re.search(r"(?:уп\.?\s*)(\d+(?:[.,]\d+)?)\s*([а-яa-z]+(?:\s*[а-яa-z]+)?)", primary_uom, re.I)
    if match:
        return format_package(match.group(1), match.group(2)), "RS PRIMARY_UOM"
    items_per_unit = num(info.get("ITEMS_PER_UNIT"))
    if items_per_unit is not None:
        unit = "шт"
        if re.search(r"\bм\b|метр", primary_uom, re.I):
            unit = "м"
        return format_package(items_per_unit, unit), "RS ITEMS_PER_UNIT"
    return "", ""


def save_checkpoint(etm_records, rs_records, counts, started):
    temp = CHECKPOINT.with_suffix(".tmp")
    payload = {
        "updated_at": time.strftime("%Y-%m-%dT%H:%M:%S%z"),
        "elapsed_seconds": round(time.monotonic() - started, 2),
        "counts": counts,
        "etm": etm_records,
        "rs": rs_records,
    }
    temp.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")
    temp.replace(CHECKPOINT)


def main():
    started = time.monotonic()
    source_rows = load_rows()
    articles = {}
    old_packages = {}
    for row in source_rows:
        article = str(row["Артикул DKC"]).strip()
        key = norm(article)
        articles.setdefault(key, article)
        if row.get("Кратность упаковки") not in (None, ""):
            old_packages[key] = row["Кратность упаковки"]

    rs_codes_by_article = load_rs_codes(articles)
    rs_code_to_articles = {}
    for article_key, codes in rs_codes_by_article.items():
        for code in codes:
            rs_code_to_articles.setdefault(code, set()).add(article_key)

    print(
        f"Подготовлено {len(articles)} артикулов; Русский Свет сопоставил "
        f"{len(rs_codes_by_article)} артикулов и {len(rs_code_to_articles)} кодов карточек.",
        flush=True,
    )

    etm_records = {}
    rs_records = {}
    counts = {"etm_done": 0, "etm_http200": 0, "etm_with_data": 0, "rs_done": 0, "rs_http200": 0}
    lock = threading.Lock()

    def fetch_etm():
        headers = {"Accept": "application/json", "User-Agent": "Mozilla/5.0"}
        http = create_etm_session()
        session_id, _ = login_etm(http, headers)
        if not session_id:
            raise RuntimeError("ETM login failed")
        last_started = 0.0
        for index, (article_key, article) in enumerate(sorted(articles.items()), start=1):
            elapsed = time.monotonic() - last_started
            if last_started and elapsed < MIN_START_GAP:
                time.sleep(MIN_START_GAP - elapsed)
            last_started = time.monotonic()
            record = {"status": None, "package": "", "source": "", "unit": ""}
            url = "https://ipro.etm.ru/api/v1/goods/" + quote(article, safe="")
            try:
                response = http.get(
                    url,
                    params={"type": "mnf", "session-id": session_id},
                    headers=headers,
                    timeout=(8, 18),
                )
                record["status"] = response.status_code
                if response.status_code == 200:
                    body = response.json()
                    data = body.get("data") if isinstance(body, dict) else None
                    status = body.get("status") if isinstance(body, dict) else None
                    if isinstance(data, dict) and (not isinstance(status, dict) or str(status.get("code")) == "200"):
                        record["unit"] = str(data.get("gdsUnitName") or "")
                        record["package"], record["source"] = choose_etm_package(data)
                        record["article"] = str(data.get("gdsArt") or "")
                        with lock:
                            counts["etm_with_data"] += 1
                    with lock:
                        counts["etm_http200"] += 1
            except Exception as exc:
                record["error"] = type(exc).__name__
            with lock:
                etm_records[article_key] = record
                counts["etm_done"] += 1
                done = counts["etm_done"]
                if done % PRINT_EVERY == 0 or done == len(articles):
                    print(f"ETM {done}/{len(articles)}; HTTP 200: {counts['etm_http200']}; карточки: {counts['etm_with_data']}", flush=True)
                if done % PRINT_EVERY == 0:
                    save_checkpoint(etm_records, rs_records, counts, started)
        http.close()

    def fetch_rs():
        http = create_rs_session()
        headers = get_rs_headers()
        last_started = 0.0
        codes = sorted(rs_code_to_articles)
        for index, code in enumerate(codes, start=1):
            elapsed = time.monotonic() - last_started
            if last_started and elapsed < MIN_START_GAP:
                time.sleep(MIN_START_GAP - elapsed)
            last_started = time.monotonic()
            record = {"status": None, "package": "", "source": ""}
            try:
                response = http.get(
                    f"{config.RS_BASE_URL}/specs/{quote(code, safe='')}",
                    headers=headers,
                    timeout=(8, 18),
                )
                record["status"] = response.status_code
                if response.status_code == 200:
                    body = response.json()
                    infos = body.get("INFO", []) if isinstance(body, dict) else []
                    info = infos[0] if isinstance(infos, list) and infos and isinstance(infos[0], dict) else {}
                    record["info"] = {k: info.get(k) for k in ("PRIMARY_UOM", "ITEMS_PER_UNIT", "MULTIPLICITY")}
                    record["package"], record["source"] = choose_rs_package(info)
                    with lock:
                        counts["rs_http200"] += 1
            except Exception as exc:
                record["error"] = type(exc).__name__
            with lock:
                for article_key in rs_code_to_articles[code]:
                    rs_records.setdefault(article_key, []).append(record)
                counts["rs_done"] += 1
                done = counts["rs_done"]
                if done % PRINT_EVERY == 0 or done == len(codes):
                    print(f"RS {done}/{len(codes)}; HTTP 200: {counts['rs_http200']}", flush=True)
                if done % PRINT_EVERY == 0:
                    save_checkpoint(etm_records, rs_records, counts, started)
        http.close()

    with ThreadPoolExecutor(max_workers=2) as executor:
        futures = [executor.submit(fetch_etm), executor.submit(fetch_rs)]
        for future in as_completed(futures):
            future.result()

    result = {
        "source_rows": source_rows,
        "articles": articles,
        "old_packages": old_packages,
        "rs_codes_by_article": {key: sorted(value) for key, value in rs_codes_by_article.items()},
        "etm": etm_records,
        "rs": rs_records,
        "counts": counts,
        "elapsed_seconds": round(time.monotonic() - started, 2),
    }
    RESULT.write_text(json.dumps(result, ensure_ascii=False), encoding="utf-8")
    save_checkpoint(etm_records, rs_records, counts, started)
    print(f"READY {len(source_rows)} rows; elapsed {result['elapsed_seconds']:.0f}s; result {RESULT}", flush=True)


if __name__ == "__main__":
    main()
